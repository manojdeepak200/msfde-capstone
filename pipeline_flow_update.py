from pptx import Presentation
from pptx.util import Inches
from pathlib import Path

base = Path(r'c:\msfde-uc007-fs-lab')
out = base / 'UC007_Claims_Processing_Capstone_Presentation.pptx'
prs = Presentation(out)

# Create a new slide at the end for the detailed pipeline
slide = prs.slides.add_slide(prs.slide_layouts[6])

# Title
textbox = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9.5), Inches(0.5))
textbox.text_frame.paragraphs[0].text = 'End-to-End Pipeline Flow'
textbox.text_frame.paragraphs[0].font.size = 24
textbox.text_frame.paragraphs[0].font.bold = True

from pptx.util import Inches

steps = [
    ('1. Claim Intake', 'Read claim folder and identify all evidence files'),
    ('2. File Classification', 'Assign each file to claim form, statement, estimate, photo, policy schedule or report'),
    ('3. Fact Extraction', 'Pull claimant, VIN, date, damage, estimate, policy and vehicle data'),
    ('4. Visual Assessment', 'Review photographs for damage severity and estimate validity'),
    ('5. Validation Rules', 'Check for missing evidence, date/VIN mismatches and suspicious patterns'),
    ('6. Policy Review', 'Compare findings against policy conditions and authority rules'),
    ('7. Recommendation', 'Return proceed, request_information or refer'),
    ('8. Human Decision', 'Handler reviews evidence and records final decision'),
]
positions = [
    (0.5, 1.4, 2.5, 1.5),
    (3.3, 1.4, 2.5, 1.5),
    (6.1, 1.4, 2.5, 1.5),
    (8.9, 1.4, 2.5, 1.5),
    (0.5, 3.4, 2.5, 1.5),
    (3.3, 3.4, 2.5, 1.5),
    (6.1, 3.4, 2.5, 1.5),
    (8.9, 3.4, 2.5, 1.5),
]

for (x, y, w, h), (title, body) in zip(positions, steps):
    box = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb = (255, 255, 255)
    box.line.color.rgb = (203, 213, 225)
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.bold = True
    p.font.size = 13
    p = tf.add_paragraph()
    p.text = body
    p.font.size = 9

prs.save(out)
print(f'Updated: {out}')
